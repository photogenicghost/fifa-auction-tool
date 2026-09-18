import csv, html, io, os, secrets, shutil, sqlite3, time, json
import base64, hashlib, re
from urllib.parse import urlsplit
from datetime import datetime, timezone
from pathlib import Path
from dotenv import load_dotenv

from fastapi import FastAPI, File, Form, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse, FileResponse
from openpyxl import load_workbook
from starlette.middleware.sessions import SessionMiddleware

BASE = Path(__file__).resolve().parent
load_dotenv(BASE / ".env")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "")
SESSION_SECRET = os.getenv("SESSION_SECRET", "")
if not ADMIN_PASSWORD or ADMIN_PASSWORD == "change-me" or len(SESSION_SECRET) < 32:
    raise RuntimeError("Set ADMIN_PASSWORD and a SESSION_SECRET of at least 32 characters in .env before starting.")
DATA_DIR = Path(os.getenv("DATA_DIR", str(BASE))).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB = DATA_DIR / "auction.db"
BACKUP_DIR = DATA_DIR / "backups"
BACKUP_DIR.mkdir(exist_ok=True)
app = FastAPI(title="FIFA Auction Tool")
app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET, same_site="lax", https_only=os.getenv("COOKIE_SECURE", "false").lower() == "true")
ONLINE_SECONDS = 35


def db():
    connection = sqlite3.connect(DB, timeout=10, isolation_level=None)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    return connection


def init():
    connection = db()
    connection.executescript("""
    CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY,v TEXT);
    INSERT OR IGNORE INTO settings VALUES('mode','test');
    INSERT OR IGNORE INTO settings VALUES('sounds','on');
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,mode TEXT,name TEXT,email TEXT,balance INTEGER,UNIQUE(mode,email));
    CREATE TABLE IF NOT EXISTS prizes(id INTEGER PRIMARY KEY,mode TEXT,name TEXT,quantity INTEGER DEFAULT 1,sku TEXT,image_url TEXT);
    CREATE TABLE IF NOT EXISTS auctions(id INTEGER PRIMARY KEY,mode TEXT,prize_id INTEGER,status TEXT,opened TEXT,closed TEXT,winner INTEGER,winning_bid INTEGER);
    CREATE UNIQUE INDEX IF NOT EXISTS one_open ON auctions(mode) WHERE status='OPEN';
    CREATE TABLE IF NOT EXISTS bids(id INTEGER PRIMARY KEY,auction_id INTEGER,user_id INTEGER,amount INTEGER,stamp TEXT,stamp_ns INTEGER);
    CREATE TABLE IF NOT EXISTS otp(id INTEGER PRIMARY KEY,email TEXT,mode TEXT,hash TEXT,expires INTEGER,used INTEGER DEFAULT 0);
    CREATE TABLE IF NOT EXISTS winners(id INTEGER PRIMARY KEY,mode TEXT,auction_id INTEGER UNIQUE,prize_name TEXT,winner_name TEXT,winning_bid INTEGER,created_at TEXT);
    CREATE TABLE IF NOT EXISTS user_presence(user_id INTEGER PRIMARY KEY,mode TEXT,last_seen REAL,login_at TEXT);
    """)
    if "ends_at" not in {row[1] for row in connection.execute("PRAGMA table_info(auctions)")}:
        connection.execute("ALTER TABLE auctions ADD COLUMN ends_at REAL")
    connection.execute("CREATE UNIQUE INDEX IF NOT EXISTS one_pending_auction ON auctions(mode) WHERE status IN ('READY','OPEN')")
    for roster_mode in ("test", "live"):
        connection.execute("INSERT OR IGNORE INTO settings(k,v) VALUES(?,?)",
                           ("session_generation_" + roster_mode, secrets.token_hex(32)))
    connection.close()


init()


def setting(key, default=""):
    connection = db()
    row = connection.execute("SELECT v FROM settings WHERE k=?", (key,)).fetchone()
    connection.close()
    return row["v"] if row else default


def mode():
    return setting("mode", "test")


def sounds_enabled():
    return setting("sounds", "on") == "on"


def e(value):
    return html.escape(str(value or ""))


def backup_db(reason="manual"):
    if not DB.exists():
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    destination = BACKUP_DIR / f"auction_{timestamp}_{reason}.db"
    source = sqlite3.connect(DB)
    target = sqlite3.connect(destination)
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()
    backups = sorted(BACKUP_DIR.glob("auction_*.db"), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in backups[50:]:
        old.unlink(missing_ok=True)
    return destination


def session_user(request, connection, current_mode):
    generation = connection.execute("SELECT v FROM settings WHERE k=?",
                                    ("session_generation_" + current_mode,)).fetchone()
    if (not generation or request.session.get("participant_mode") != current_mode
            or request.session.get("participant_generation") != generation[0]):
        return None
    return connection.execute("SELECT * FROM users WHERE id=? AND mode=?",
                              (request.session.get("uid"), current_mode)).fetchone()


def invalidate_participants(connection, current_mode):
    connection.execute("UPDATE settings SET v=? WHERE k=?",
                       (secrets.token_hex(32), "session_generation_" + current_mode))
    connection.execute("DELETE FROM user_presence WHERE mode=?", (current_mode,))


def touch_user(request):
    connection = db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        current_mode = connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        user = session_user(request, connection, current_mode)
        if user:
            connection.execute("""INSERT INTO user_presence(user_id,mode,last_seen,login_at)
                VALUES(?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET last_seen=excluded.last_seen""",
                (user["id"], current_mode, time.time(), datetime.now(timezone.utc).isoformat()))
        connection.execute("COMMIT")
        return bool(user)
    finally:
        connection.close()


def relative_seen(seconds):
    age = max(0, int(time.time() - seconds))
    if age < 5:
        return "just now"
    if age < 60:
        return f"{age}s ago"
    if age < 3600:
        return f"{age // 60}m ago"
    return f"{age // 3600}h ago"


def auction_clock_script():
    return """<script>
    let auctionClock=null,clockUntil=0;
    function auctionIsOpen(a){return !!(a&&a.status==='OPEN'&&a.ends_at!=null&&auctionClock?.id===a.id&&performance.now()<clockUntil)}
    function auctionStatus(a){
      if(!a)return 'Waiting for the next prize';
      if(a.status==='READY')return 'Item loaded - waiting to start';
      if(a.status==='CANCELLED')return 'Auction cancelled';
      if(a.status==='ENDED'||(a.status==='OPEN'&&!auctionIsOpen(a)))return 'Bidding closed - awaiting award';
      return a.status==='OPEN'?'Bidding open':'Auction closed';
    }
    function updateAuctionClock(s){
      auctionClock=s.auction;
      clockUntil=performance.now()+Math.max(0,((s.auction?.ends_at||0)-s.server_now)*1000);
      tickAuctionClock();
    }
    function tickAuctionClock(){
      const a=auctionClock,t=document.getElementById('countdown');if(!t)return;
      const open=auctionIsOpen(a),pending=a&&(a.status==='OPEN'||a.status==='ENDED');
      t.textContent=a?.status==='READY'?'30 seconds - ready':open?Math.ceil(Math.max(0,clockUntil-performance.now())/1000)+' seconds':pending?'0 seconds - bidding closed':'';
      const status=document.getElementById('auction-status')||document.getElementById('admin-auction-status');
      if(status)status.textContent=auctionStatus(a);
      const displayStatus=document.getElementById('status');
      if(displayStatus&&pending&&!open){displayStatus.textContent=auctionStatus(a);displayStatus.className='closed'}
      if(!open)for(const id of ['submit','amount','next']){const b=document.getElementById(id);if(b)b.disabled=true}
      const endEarly=document.getElementById('end-bidding-early');if(endEarly)endEarly.disabled=!open;
      const award=document.getElementById('review-award');if(award)award.disabled=!pending||open;
    }
    setInterval(tickAuctionClock,100);
    </script>"""


def auction_controls(options):
    a=current();pending=a and a['status'] in ('READY','OPEN')
    aid=a['id'] if pending else 0
    return f'''<div class="card"><h2>Auction Controls</h2>
    <h3>{e(a['prize']) if pending else 'Select a prize to load'}</h3>
    <p id="admin-auction-status"></p><div id="countdown" class="leader" role="timer"></div>
    <form method="post" action="/admin/load"><select name="prize_id" required aria-label="Prize to load">{options}</select><button {'disabled' if pending else ''}>Load Item</button></form>
    <form method="post" action="/admin/start"><input type="hidden" name="auction_id" value="{aid}"><button class="green" {'disabled' if not pending or a['status']!='READY' else ''}>Start 30-Second Auction</button></form>
    <form method="post" action="/admin/end" onsubmit="return confirm('End bidding now? The highest bid will be kept for review and award.')"><input type="hidden" name="auction_id" value="{aid}"><button id="end-bidding-early" class="danger" disabled>End Bidding Early</button></form>
    <form method="post" action="/admin/close/preview"><button id="review-award" class="gold" disabled>Review and Award</button></form>
    <form method="post" action="/admin/cancel" onsubmit="return confirm('Cancel this auction? No points or stock will be deducted.')"><input type="hidden" name="auction_id" value="{aid}"><button class="danger" {'disabled' if not pending else ''}>Cancel Auction</button></form></div>'''


def page(body):
    css = """<style>
    *{box-sizing:border-box}body{margin:0;font-family:Arial;background:#f3f6fa;color:#14233a}
    nav{display:flex;gap:18px;flex-wrap:wrap;background:#071a33;padding:16px;color:white}
    nav b{margin-right:auto}nav a{color:white;text-decoration:none;font-weight:bold}
    main{max-width:1150px;margin:25px auto;padding:0 16px}.card{background:white;padding:20px;border-radius:14px;margin:15px 0;box-shadow:0 4px 18px #1232}
    .grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(270px,1fr));gap:15px}.summary{grid-template-columns:repeat(auto-fit,minmax(150px,1fr))}
    button,input,select{padding:12px;margin:5px;border-radius:8px;border:1px solid #abc;font:inherit}
    button{background:#0868d7;color:white;border:0;font-weight:bold;cursor:pointer}.danger{background:#bd2424}.green{background:#17813e}.gold{background:#f5c542;color:#111}.secondary{background:#65758b}
    button:disabled{opacity:.45;cursor:not-allowed}#countdown{font-variant-numeric:tabular-nums;margin:12px 0;color:#0868d7}.big{font-size:clamp(2rem,8vw,5rem);font-weight:bold}.leader{font-size:clamp(1.5rem,5vw,3.5rem)}.metric{font-size:2rem;font-weight:bold}
    table{width:100%;border-collapse:collapse}td,th{padding:9px;border-bottom:1px solid #ddd;text-align:left}
    .display{text-align:center}.open{color:#17813e}.closed{color:#bd2424}.prize{width:100%;height:220px;object-fit:contain}
    .trophy{font-size:clamp(2.5rem,8vw,6rem);font-weight:bold;color:#b8860b}.winner{background:linear-gradient(135deg,#fff8d8,#ffffff);padding:28px;border-radius:18px}
    .warning{background:#fff4d7;border-left:5px solid #e3a008;padding:12px;margin:12px 0}.online-dot{display:inline-block;width:10px;height:10px;border-radius:50%;background:#19a350;margin-right:8px}
    .bid-flash{animation:flash 1.2s ease}@keyframes flash{0%,100%{transform:scale(1)}30%{transform:scale(1.08);color:#0868d7}}
    #confetti{position:fixed;inset:0;pointer-events:none;z-index:9999}
    @media(max-width:650px){table{font-size:.84rem}td,th{padding:6px}button,input,select{max-width:100%}}
    </style>"""
    nav = '<nav><b>FIFA Auction Tool</b><a href="/">Login</a><a href="/auction">Auction</a><a href="/catalog">Prizes</a><a href="/display">Display</a><a href="/admin">Admin</a></nav>'
    return HTMLResponse(f'<!doctype html><meta name="viewport" content="width=device-width"><title>FIFA Auction Tool</title>{css}{auction_clock_script()}{nav}<main>{body}</main>')


def current(connection=None, current_mode=None):
    owns_connection = connection is None
    connection = connection if connection is not None else db()
    row = connection.execute("""SELECT a.*,p.name prize,p.image_url,
        (SELECT amount FROM bids WHERE auction_id=a.id ORDER BY amount DESC,stamp_ns,id LIMIT 1) high,
        (SELECT u.name FROM bids b JOIN users u ON u.id=b.user_id WHERE b.auction_id=a.id ORDER BY b.amount DESC,b.stamp_ns,b.id LIMIT 1) leader,
        (SELECT id FROM bids WHERE auction_id=a.id ORDER BY amount DESC,stamp_ns,id LIMIT 1) high_bid_id
        FROM auctions a JOIN prizes p ON p.id=a.prize_id WHERE a.mode=?
        ORDER BY CASE WHEN a.status IN ('READY','OPEN') THEN 0 ELSE 1 END,a.id DESC LIMIT 1""", (current_mode if current_mode is not None else mode(),)).fetchone()
    if owns_connection: connection.close()
    return row


def auction_payload(row, now=None):
    if not row:
        return None
    result = dict(row)
    result['image_url'] = normalize_image_url(result.get('image_url'))
    now = time.time() if now is None else now
    if result['status'] == 'OPEN' and (result['ends_at'] is None or now >= result['ends_at']):
        result['status'] = 'ENDED'
    return result


class Hub:
    def __init__(self):
        self.clients = []

    async def add(self, websocket):
        await websocket.accept()
        self.clients.append(websocket)

    async def push(self):
        for websocket in self.clients[:]:
            try:
                await websocket.send_text("update")
            except Exception:
                self.clients.remove(websocket)


hub = Hub()


@app.get("/healthz")
def healthcheck():
    connection = db()
    try:
        connection.execute("SELECT 1 FROM settings LIMIT 1").fetchone()
    finally:
        connection.close()
    return {"ok": True}


def participant_script(user_id):
    return f'''<script>
    const wsProto=location.protocol==='https:'?'wss':'ws';
    {live_updates_js()}
    function heartbeat(){{fetch('/heartbeat',{{method:'POST',credentials:'same-origin'}}).catch(()=>{{}})}}
    heartbeat();setInterval(heartbeat,10000);
    </script>'''


def admin_script():
    return '''<script>
    const wsProto=location.protocol==='https:'?'wss':'ws';
    ''' + live_updates_js() + '''
    async function refreshClock(){try{const r=await fetch('/api/display',{cache:'no-store'});if(r.ok)updateAuctionClock(await r.json())}catch(e){}}
    refreshClock();setInterval(refreshClock,1000);
    </script>'''


def live_updates_js():
    return '''
    let socket, retryTimer, connectedOnce=false;
    function connectUpdates(){
      clearTimeout(retryTimer);
      socket=new WebSocket(`${wsProto}://${location.host}/ws`);
      socket.onopen=()=>{if(connectedOnce)location.reload();connectedOnce=true};
      socket.onmessage=()=>location.reload();
      socket.onclose=()=>{retryTimer=setTimeout(connectUpdates,3000)};
      socket.onerror=()=>socket.close();
    }
    connectUpdates();
    window.addEventListener('online',()=>location.reload());
    document.addEventListener('visibilitychange',()=>{if(!document.hidden)location.reload()});
    '''


def display_script(auction_id, high_bid_id, winner=False):
    sound = "true" if sounds_enabled() else "false"
    return f'''<canvas id="confetti"></canvas><script>
    const soundOn={sound};const auctionId={auction_id or 0};const highBidId={high_bid_id or 0};
    function tone(freq,duration,type='sine'){{if(!soundOn)return;try{{const c=new(window.AudioContext||window.webkitAudioContext)();const o=c.createOscillator();const g=c.createGain();o.type=type;o.frequency.value=freq;o.connect(g);g.connect(c.destination);g.gain.setValueAtTime(.18,c.currentTime);g.gain.exponentialRampToValueAtTime(.001,c.currentTime+duration);o.start();o.stop(c.currentTime+duration)}}catch(e){{}}}}
    const lastA=Number(localStorage.getItem('auctionId')||0),lastB=Number(localStorage.getItem('highBidId')||0);
    if(auctionId===lastA && highBidId>lastB){{tone(880,.18);setTimeout(()=>tone(1175,.18),150)}}
    localStorage.setItem('auctionId',auctionId);localStorage.setItem('highBidId',highBidId);
    function confetti(){{const c=document.getElementById('confetti'),x=c.getContext('2d');c.width=innerWidth;c.height=innerHeight;let p=Array.from({{length:180}},()=>({{x:Math.random()*c.width,y:-20-Math.random()*c.height,vx:(Math.random()-.5)*5,vy:2+Math.random()*5,r:3+Math.random()*5,h:Math.random()*360}}));let n=0;function f(){{x.clearRect(0,0,c.width,c.height);p.forEach(q=>{{q.x+=q.vx;q.y+=q.vy;q.vy+=.03;x.fillStyle=`hsl(${{q.h}} 85% 55%)`;x.fillRect(q.x,q.y,q.r,q.r)}});if(n++<240)requestAnimationFrame(f)}}f()}}
    if({str(winner).lower()}){{confetti();tone(523,.25);setTimeout(()=>tone(659,.25),250);setTimeout(()=>tone(784,.45),500)}}
    const wsProto=location.protocol==='https:'?'wss':'ws';{live_updates_js()}
    </script>'''


def normalize_image_url(value):
    """Accept browser image sources without restricting hosts or file extensions."""
    value = str(value or '').strip()
    if value.startswith('//'):
        value = 'https:' + value
    if value.lower().startswith('www.'):
        value = 'https://' + value
    if re.match(r'^data:image/(?:png|jpeg|jpg|gif|webp|avif|bmp|x-icon);base64,', value, re.I):
        try:
            header, encoded = value.split(',', 1)
            encoded = re.sub(r'\s+', '', encoded)
            if not encoded or len(encoded) > 4_000_000:
                return ''
            base64.b64decode(encoded, validate=True)
            return header.lower() + ',' + encoded
        except ValueError:
            return ''
    try:
        parsed = urlsplit(value)
        if parsed.scheme.lower() in ('http', 'https') and parsed.hostname and not parsed.username and not parsed.password:
            return value
    except ValueError:
        pass
    return ''


def image_cell_source(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    value = str(cell.value or '').strip()
    # Read literal links in Excel formulas without executing workbook formulas.
    match = re.match(r'^=\s*(?:_xlfn\.)?(?:HYPERLINK|IMAGE)\s*\(\s*"((?:[^"]|"")*)"', value, re.I)
    return match.group(1).replace('""', '"') if match else value


def parse_xlsx(data):
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    warnings, users, prizes = [], [], []
    sheet = workbook["Total Points"] if "Total Points" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header = None
    for index, row in enumerate(rows[:20]):
        values = [str(value or "").strip().lower() for value in row]
        if any("total points" in value or value == "points" for value in values) and "name" in values and any("email" in value for value in values):
            header = index; points_col = next(i for i,v in enumerate(values) if "total points" in v or v=="points"); name_col=values.index("name"); email_col=next(i for i,v in enumerate(values) if "email" in v); break
    if header is None: raise ValueError("Participant headers Total Points, NAME, and Email were not found.")
    seen=set()
    for row_number,row in enumerate(rows[header+1:],header+2):
        name=str(row[name_col] or "").strip() if len(row)>name_col else ""; email=str(row[email_col] or "").strip().lower() if len(row)>email_col else ""; raw=row[points_col] if len(row)>points_col else None
        if not name and not email: continue
        try: points=max(0,int(round(float(raw))))
        except Exception: warnings.append(f"Participant row {row_number} skipped: invalid points"); continue
        if not name or "@" not in email: warnings.append(f"Participant row {row_number} skipped: missing name/email"); continue
        if email in seen: warnings.append(f"Duplicate email skipped: {email}"); continue
        seen.add(email); users.append((name,email,points))
    if "Auction Prize List" in workbook.sheetnames:
        source_book = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
        source_sheet = source_book['Auction Prize List']
        sheet=workbook["Auction Prize List"]; rows=list(sheet.iter_rows(values_only=True)); header=None
        for index,row in enumerate(rows[:20]):
            values=[str(value or "").strip().lower() for value in row]
            if any("product description" in v for v in values) and "sku" in values:
                header=index; name_col=next(i for i,v in enumerate(values) if "product description" in v); sku_col=values.index("sku"); image_col=next((i for i,v in enumerate(values) if "image" in v),None); break
        if header is None: warnings.append("Prize sheet headers not recognized.")
        else:
            grouped={}
            for row_number,row in enumerate(rows[header+1:],header+2):
                name=str(row[name_col] or "").strip() if len(row)>name_col else ""; sku=str(row[sku_col] or "").strip() if len(row)>sku_col else ""
                if not name: continue
                if not sku:
                    sku = 'AUTO-' + hashlib.sha256(name.casefold().encode()).hexdigest()[:20]
                raw_image = image_cell_source(source_sheet.cell(row_number,image_col+1)) if image_col is not None else ''
                image = normalize_image_url(raw_image)
                if raw_image and not image:
                    warnings.append(f'Prize row {row_number}: use a public HTTP(S) image URL or a base64 image link.')
                key=sku.lower(); grouped.setdefault(key,[name,sku,image,0]); grouped[key][3]+=1
                if image and not grouped[key][2]: grouped[key][2]=image
            prizes=list(grouped.values())
        source_book.close()
    workbook.close()
    return users,prizes,warnings


@app.get("/")
def login():
    return page('''<div class="card"><h1>Participant Login</h1><div class="warning">Login is based on the honor system. Only use your own registered work email address.</div><form method="post" action="/login"><input type="email" name="email" placeholder="Registered work email" required><button>Continue to Auction</button></form></div>''')


@app.post("/login")
def participant_login(request:Request,email:str=Form(...)):
    connection = db()
    try:
        connection.execute("BEGIN")
        m = connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        user = connection.execute("SELECT id FROM users WHERE mode=? AND lower(email)=?",
                                  (m, email.lower().strip())).fetchone()
        generation = connection.execute("SELECT v FROM settings WHERE k=?",
                                        ("session_generation_" + m,)).fetchone()[0]
    finally:
        connection.close()
    if not user: return page('<div class="card"><h2>Email not registered</h2><p>Use the email address included in the imported participant list.</p><a href="/">Try again</a></div>')
    request.session.clear()
    request.session.update(uid=user["id"], participant_mode=m, participant_generation=generation)
    touch_user(request)
    return RedirectResponse("/auction",303)


@app.post("/heartbeat")
def heartbeat(request:Request):
    return JSONResponse({"ok":touch_user(request)})


@app.get("/logout")
def logout(request:Request):
    connection = db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        m = connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        user = session_user(request, connection, m)
        if user:
            connection.execute("DELETE FROM user_presence WHERE user_id=?", (user["id"],))
        connection.execute("COMMIT")
    finally:
        connection.close()
    request.session.clear(); return RedirectResponse("/",303)


@app.get("/auction")
def auction(request:Request):
    if not request.session.get("uid"):return RedirectResponse("/",303)
    return participant_page(request)



@app.post("/bid")
async def bid(request:Request,amount:int=Form(...),auction_id:int=Form(0)):
    uid=request.session.get("uid");connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE");m=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0];user=session_user(request,connection,m);active=connection.execute("SELECT * FROM auctions WHERE mode=? AND status='OPEN'",(m,)).fetchone();high=connection.execute("SELECT amount FROM bids WHERE auction_id=? ORDER BY amount DESC,stamp_ns,id LIMIT 1",(active["id"],)).fetchone() if active else None
        if not user:raise ValueError("Sign in with your registered email again.")
        if not active:raise ValueError("Bidding is closed. Wait for the next prize.")
        if active["ends_at"] is None or time.time() >= active["ends_at"]:
            raise ValueError("Time is up. Bidding is closed.")
        if auction_id!=active["id"]:raise ValueError("The prize changed. Refresh and review the prize before bidding.")
        if amount<1:raise ValueError("Enter at least 1 point.")
        if amount>user["balance"]:raise ValueError("This bid exceeds your available balance.")
        if high and amount<=high["amount"]:raise ValueError(f"Another bid arrived first. Bid at least {high['amount']+1} points.")
        connection.execute("INSERT INTO bids(auction_id,user_id,amount,stamp,stamp_ns) VALUES(?,?,?,?,?)",(active["id"],uid,amount,datetime.now(timezone.utc).isoformat(),time.time_ns()));connection.execute("COMMIT")
    except Exception as error:
        try:connection.execute("ROLLBACK")
        except Exception:pass
        connection.close()
        if request.headers.get("accept")=="application/json":return JSONResponse({"ok":False,"message":str(error) if isinstance(error,ValueError) else "Bid could not be submitted. Please try again."},status_code=400)
        return page('<div class="card"><h2>Bid rejected</h2><p>Check the auction status, current high bid, and available balance.</p><a href="/auction">Return</a></div>')
    connection.close();await hub.push()
    if request.headers.get("accept")=="application/json":return JSONResponse({"ok":True,"message":"Bid accepted."})
    return RedirectResponse("/auction",303)


@app.get("/catalog")
def catalog(request:Request):
    touch_user(request);connection=db();prizes=connection.execute("SELECT * FROM prizes WHERE mode=? ORDER BY name",(mode(),)).fetchall();connection.close();cards=[]
    for prize in prizes:
        source=normalize_image_url(prize['image_url'])
        image=f'<img class="prize" src="{e(source)}" alt="{e(prize["name"])}" loading="lazy" referrerpolicy="no-referrer" onerror="this.hidden=true;this.nextElementSibling.hidden=false"><p hidden>Image unavailable. Use a direct public image link.</p>' if source else ''
        cards.append(f'<div class="card">{image}<h2>{e(prize["name"])}</h2><p>Quantity: {prize["quantity"]}</p></div>')
    script=participant_script(request.session.get("uid")) if request.session.get("uid") else ""
    return page('<h1>Prize Catalog</h1><div class="grid">'+"".join(cards)+"</div>"+script)


@app.get("/display")
def display():
    return display_page()



@app.get("/admin")
def admin(request:Request):
    if not request.session.get("admin"):return page('<div class="card"><h1>Admin Login</h1><form method="post" action="/admin/login"><input type="password" name="password" required><button>Sign in</button></form></div>')
    m=mode();cutoff=time.time()-ONLINE_SECONDS;connection=db();users=connection.execute("SELECT * FROM users WHERE mode=? ORDER BY name",(m,)).fetchall();prizes=connection.execute("SELECT * FROM prizes WHERE mode=? AND quantity>0 ORDER BY name",(m,)).fetchall();bids=connection.execute("SELECT b.stamp,u.name,p.name prize,b.amount FROM bids b JOIN users u ON u.id=b.user_id JOIN auctions a ON a.id=b.auction_id JOIN prizes p ON p.id=a.prize_id WHERE a.mode=? ORDER BY b.id DESC LIMIT 50",(m,)).fetchall();winners=connection.execute("SELECT * FROM winners WHERE mode=? ORDER BY id DESC LIMIT 20",(m,)).fetchall();online=connection.execute("SELECT u.name,u.email,p.last_seen FROM user_presence p JOIN users u ON u.id=p.user_id WHERE p.mode=? AND p.last_seen>=? ORDER BY p.last_seen DESC",(m,cutoff)).fetchall();total_bids=connection.execute("SELECT COUNT(*) n FROM bids b JOIN auctions a ON a.id=b.auction_id WHERE a.mode=?",(m,)).fetchone()["n"];closed=connection.execute("SELECT COUNT(*) n FROM auctions WHERE mode=? AND status='CLOSED'",(m,)).fetchone()["n"];remaining=connection.execute("SELECT COALESCE(SUM(quantity),0) n FROM prizes WHERE mode=?",(m,)).fetchone()["n"];connection.close()
    options="".join(f'<option value="{p["id"]}">{e(p["name"])} ({p["quantity"]})</option>' for p in prizes);user_rows="".join(f'<tr><td>{e(u["name"])}</td><td>{e(u["email"])}</td><td><form method="post" action="/admin/user/update"><input type="hidden" name="user_id" value="{u["id"]}"><input type="number" name="balance" min="0" value="{u["balance"]}" style="width:100px"><button>Save</button></form></td></tr>' for u in users);bid_rows="".join(f'<tr><td>{e(x["stamp"])}</td><td>{e(x["name"])}</td><td>{e(x["prize"])}</td><td>{x["amount"]}</td></tr>' for x in bids);winner_rows="".join(f'<tr><td>{e(x["prize_name"])}</td><td>{e(x["winner_name"])}</td><td>{x["winning_bid"]}</td><td>{e(x["created_at"])}</td></tr>' for x in winners);online_rows="".join(f'<tr><td><span class="online-dot"></span>{e(x["name"])}</td><td>{e(x["email"])}</td><td>{relative_seen(x["last_seen"])}</td></tr>' for x in online)
    reset='<form method="post" action="/admin/reset-test" onsubmit="return confirm(\'Reset all TEST data?\')"><button class="danger">Reset TEST Mode</button></form>' if m=="test" else "";sound_label="Disable Sounds" if sounds_enabled() else "Enable Sounds"
    body=f'''<h1>Admin: {m.upper()}</h1><div class="grid summary"><div class="card"><div>Connected</div><div class="metric">{len(online)}</div><small>of {len(users)}</small></div><div class="card"><div>Total Bids</div><div class="metric">{total_bids}</div></div><div class="card"><div>Auctions Closed</div><div class="metric">{closed}</div></div><div class="card"><div>Remaining Prizes</div><div class="metric">{remaining}</div></div></div><div class="card"><form method="post" action="/admin/mode"><button class="gold">Switch Test/Live</button></form>{reset}<form method="post" action="/admin/sounds"><button class="secondary">{sound_label}</button></form><form method="post" action="/admin/backup"><button class="secondary">Download Backup</button></form></div><div class="card"><h2>Connected Participants ({len(online)})</h2><table><tr><th>Name</th><th>Email</th><th>Last Activity</th></tr>{online_rows}</table></div><div class="card"><h2>Import Excel Workbook</h2><form method="post" action="/admin/import/preview" enctype="multipart/form-data"><input type="file" name="workbook" accept=".xlsx" required><select name="action"><option value="replace">Replace current mode data</option><option value="append">Append/update names; keep balances and stock</option></select><button>Preview Import</button></form><a href="/admin/export/users">Participants CSV</a> | <a href="/admin/export/results">Results CSV</a></div>{auction_controls(options)}<div class="card"><h2>Participants and Balance Editing</h2><table><tr><th>Name</th><th>Email</th><th>Balance</th></tr>{user_rows}</table></div><div class="card"><h2>Bid History</h2><table><tr><th>UTC Time</th><th>Bidder</th><th>Prize</th><th>Amount</th></tr>{bid_rows}</table></div><div class="card"><h2>Winner History</h2><table><tr><th>Prize</th><th>Winner</th><th>Winning Bid</th><th>UTC Time</th></tr>{winner_rows}</table></div>'''
    return page(body+admin_script())


@app.post("/admin/login")
def admin_login(request:Request,password:str=Form(...)):
    if secrets.compare_digest(password.encode("utf-8"),ADMIN_PASSWORD.encode("utf-8")):request.session.clear();request.session["admin"]=True
    return RedirectResponse("/admin",303)


async def import_workbook(request:Request,workbook:UploadFile=File(...),action:str=Form("replace"),confirmed:bool=False,expected_mode:str=""):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    if not confirmed:return page('<div class="card"><h2>Preview the workbook before importing.</h2><a href="/admin">Return</a></div>')
    try:users,prizes,warnings=parse_xlsx(await workbook.read())
    except Exception as error:return page(f'<div class="card"><h2>Import failed</h2><p>{e(error)}</p><a href="/admin">Return</a></div>')
    backup_db("before_import");m=mode();connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        actual_mode=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        if actual_mode!=expected_mode or m!=expected_mode:raise ValueError("Mode changed. Preview the workbook again.")
        if action=="replace":
            invalidate_participants(connection, m)
            if connection.execute("SELECT 1 FROM auctions WHERE mode=? AND status IN ('READY','OPEN')",(m,)).fetchone():raise ValueError("Close the open auction or cancel the loaded item before replacing data.")
            connection.execute("DELETE FROM bids WHERE auction_id IN (SELECT id FROM auctions WHERE mode=?)",(m,));connection.execute("DELETE FROM winners WHERE mode=?",(m,));connection.execute("DELETE FROM auctions WHERE mode=?",(m,));connection.execute("DELETE FROM otp WHERE mode=?",(m,));connection.execute("DELETE FROM user_presence WHERE mode=?",(m,));connection.execute("DELETE FROM users WHERE mode=?",(m,));connection.execute("DELETE FROM prizes WHERE mode=?",(m,))
        for name,email,points in users:connection.execute("INSERT INTO users(mode,name,email,balance) VALUES(?,?,?,?) ON CONFLICT(mode,email) DO UPDATE SET name=excluded.name",(m,name,email,points))
        for name,sku,image,quantity in prizes:
            existing=connection.execute("SELECT id FROM prizes WHERE mode=? AND sku=?",(m,sku)).fetchone()
            if existing:connection.execute("UPDATE prizes SET name=?,image_url=? WHERE id=?",(name,image,existing["id"]))
            else:connection.execute("INSERT INTO prizes(mode,name,quantity,sku,image_url) VALUES(?,?,?,?,?)",(m,name,quantity,sku,image))
        connection.execute("COMMIT")
    except Exception as error:
        connection.execute("ROLLBACK");connection.close();return page(f'<div class="card"><h2>Import failed</h2><p>{e(error)}</p><a href="/admin">Return</a></div>')
    connection.close();backup_db("after_import");warning_list="".join(f"<li>{e(x)}</li>" for x in warnings);return page(f'<div class="card"><h1>Import complete</h1><p>{len(users)} participants and {len(prizes)} unique prizes imported into {m.upper()} mode.</p>{("<ul>"+warning_list+"</ul>") if warning_list else ""}<a href="/admin">Return</a></div>')


@app.post("/admin/user/update")
async def update_user(request:Request,user_id:int=Form(...),balance:int=Form(...)):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    if balance<0:return page('<div class="card"><h2>Balance cannot be negative.</h2><a href="/admin">Return</a></div>')
    connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        m=mode()
        outstanding=connection.execute("SELECT MAX(b.amount) FROM bids b JOIN auctions a ON a.id=b.auction_id WHERE b.user_id=? AND a.mode=? AND a.status='OPEN'",(user_id,m)).fetchone()[0]
        if outstanding and balance<outstanding:raise ValueError("Balance cannot be lower than this participant's bid in the open auction.")
        connection.execute("UPDATE users SET balance=? WHERE id=? AND mode=?",(balance,user_id,m));connection.execute("COMMIT")
    except Exception as error:
        connection.execute("ROLLBACK");return page(f'<div class="card"><h2>Balance update failed</h2><p>{e(error)}</p><a href="/admin">Return</a></div>')
    finally:connection.close()
    await hub.push();return RedirectResponse("/admin",303)


@app.post("/admin/reset-test")
async def reset_test(request:Request):
    if not request.session.get("admin") or mode()!="test":return RedirectResponse("/admin",303)
    backup_db("before_test_reset");connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        if connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0] != "test":
            raise ValueError("Mode changed. Return to TEST mode before resetting.")
        invalidate_participants(connection, "test")
        connection.execute("DELETE FROM bids WHERE auction_id IN (SELECT id FROM auctions WHERE mode='test')");connection.execute("DELETE FROM winners WHERE mode='test'");connection.execute("DELETE FROM auctions WHERE mode='test'");connection.execute("DELETE FROM otp WHERE mode='test'");connection.execute("DELETE FROM user_presence WHERE mode='test'");connection.execute("DELETE FROM users WHERE mode='test'");connection.execute("DELETE FROM prizes WHERE mode='test'");connection.execute("COMMIT")
    except Exception as error:
        connection.execute("ROLLBACK")
        return page(f'<h2>Test reset failed</h2><p>{e(error)}</p><a href="/admin">Return</a>')
    finally:connection.close()
    await hub.push();return RedirectResponse("/admin",303)


@app.post("/admin/open")
@app.post("/admin/load")
async def load_auction(request:Request,prize_id:int=Form(...)):
    if not request.session.get("admin"):
        return RedirectResponse("/admin",303)
    connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        m=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        prize=connection.execute("SELECT 1 FROM prizes WHERE id=? AND mode=? AND quantity>0",(prize_id,m)).fetchone()
        if not prize:raise ValueError("Choose an available prize in the current mode.")
        if connection.execute("SELECT 1 FROM auctions WHERE mode=? AND status IN ('READY','OPEN')",(m,)).fetchone():
            raise ValueError("Award or cancel the current item before loading another.")
        connection.execute("INSERT INTO auctions(mode,prize_id,status) VALUES(?,?,'READY')",(m,prize_id))
        connection.execute("COMMIT")
    except Exception as error:
        connection.execute("ROLLBACK")
        return page(f'<h2>Could not load item</h2><p>{e(error)}</p><a href="/admin">Return</a>')
    finally:connection.close()
    await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/start")
async def start_auction(request:Request,auction_id:int=Form(...)):
    if not request.session.get("admin"):
        return RedirectResponse("/admin",303)
    connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        m=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        active=connection.execute("SELECT * FROM auctions WHERE mode=? AND status IN ('READY','OPEN')",(m,)).fetchone()
        if not active or active['id']!=auction_id:raise ValueError("The item changed. Review the loaded item before starting.")
        if active['status']!='READY':raise ValueError("This auction has already started. Its timer cannot be restarted.")
        now=time.time()
        connection.execute("UPDATE auctions SET status='OPEN',opened=?,ends_at=? WHERE id=?",
                           (datetime.fromtimestamp(now,timezone.utc).isoformat(),now+30,auction_id))
        connection.execute("COMMIT")
    except Exception as error:
        connection.execute("ROLLBACK")
        return page(f'<h2>Could not start auction</h2><p>{e(error)}</p><a href="/admin">Return</a>')
    finally:connection.close()
    await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/end")
async def end_bidding_early(request:Request,auction_id:int=Form(...)):
    if not request.session.get("admin"):
        return RedirectResponse("/admin",303)
    connection=db()
    try:
        # Serialize with bids so the cutoff and accepted bid history agree.
        connection.execute("BEGIN IMMEDIATE")
        m=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        active=connection.execute("SELECT * FROM auctions WHERE mode=? AND status IN ('READY','OPEN')",(m,)).fetchone()
        if not active or active['id']!=auction_id:
            raise ValueError("The auction changed. Review the current item before ending bidding.")
        if active['status']!='OPEN':
            raise ValueError("Start the auction before ending bidding, or cancel the loaded item.")
        now=time.time()
        if active['ends_at'] is not None and active['ends_at']>now:
            connection.execute("UPDATE auctions SET ends_at=? WHERE id=?",(now,auction_id))
        connection.execute("COMMIT")
    except Exception as error:
        connection.execute("ROLLBACK")
        return page(f'<h2>Could not end bidding</h2><p>{e(error)}</p><a href="/admin">Return</a>')
    finally:connection.close()
    await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/cancel")
async def cancel_auction(request:Request,auction_id:int=Form(...)):
    if not request.session.get("admin"):
        return RedirectResponse("/admin",303)
    connection=db()
    try:
        connection.execute("BEGIN IMMEDIATE")
        m=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        connection.execute("UPDATE auctions SET status='CANCELLED',closed=? WHERE id=? AND mode=? AND status IN ('READY','OPEN')",
                           (datetime.now(timezone.utc).isoformat(),auction_id,m))
        connection.execute("COMMIT")
    finally:connection.close()
    await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/close")
async def close_auction(request:Request,auction_id:int=Form(0),bid_id:int=Form(-1)):
    if request.session.get("admin"):
        m=mode();connection=db();success=False
        try:
            connection.execute("BEGIN IMMEDIATE");active=connection.execute("SELECT * FROM auctions WHERE mode=? AND status='OPEN'",(m,)).fetchone()
            if active:
                if active['ends_at'] is not None and time.time() < active['ends_at']:
                    raise ValueError("Wait until the timer ends or use End Bidding Early before awarding the prize.")
                if auction_id!=active["id"]:raise ValueError("The auction changed. Review the award again.")
                winning=connection.execute("SELECT * FROM bids WHERE auction_id=? ORDER BY amount DESC,stamp_ns,id LIMIT 1",(active["id"],)).fetchone();closed=datetime.now(timezone.utc).isoformat()
                if bid_id!=(winning["id"] if winning else 0):raise ValueError("A new bid arrived. Review the latest winner before confirming.")
                if winning:
                    user=connection.execute("SELECT name,balance FROM users WHERE id=? AND mode=?",(winning["user_id"],m)).fetchone();prize=connection.execute("SELECT name,quantity FROM prizes WHERE id=? AND mode=?",(active["prize_id"],m)).fetchone()
                    if not user or user["balance"]<winning["amount"]:raise ValueError("Winner balance is no longer sufficient.")
                    if not prize or prize["quantity"]<1:raise ValueError("Prize is no longer available.")
                    connection.execute("UPDATE users SET balance=balance-? WHERE id=?",(winning["amount"],winning["user_id"]));connection.execute("UPDATE prizes SET quantity=quantity-1 WHERE id=?",(active["prize_id"],));connection.execute("UPDATE auctions SET status='CLOSED',closed=?,winner=?,winning_bid=? WHERE id=?",(closed,winning["user_id"],winning["amount"],active["id"]));connection.execute("INSERT OR REPLACE INTO winners(mode,auction_id,prize_name,winner_name,winning_bid,created_at) VALUES(?,?,?,?,?,?)",(m,active["id"],prize["name"],user["name"],winning["amount"],closed))
                else:connection.execute("UPDATE auctions SET status='CLOSED',closed=? WHERE id=?",(closed,active["id"]))
            connection.execute("COMMIT");success=True
        except Exception as error:
            try:connection.execute("ROLLBACK")
            except Exception:pass
            connection.close()
            return page(f'<div class="card"><h2>Could not close auction</h2><p>{e(error)}</p><a href="/admin">Return</a></div>')
        connection.close()
        if success:backup_db("auction_closed")
        await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/mode")
async def switch_mode(request:Request):
    if request.session.get("admin"):
        connection=db()
        try:
            connection.execute("BEGIN IMMEDIATE")
            if connection.execute("SELECT 1 FROM auctions WHERE status IN ('READY','OPEN')").fetchone():raise ValueError("Award or cancel the current item before switching modes.")
            old=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
            invalidate_participants(connection, "test")
            invalidate_participants(connection, "live")
            connection.execute("UPDATE settings SET v=? WHERE k='mode'",("live" if old=="test" else "test",));connection.execute("COMMIT")
        except Exception as error:
            connection.execute("ROLLBACK");return page(f'<div class="card"><h2>Could not switch mode</h2><p>{e(error)}</p><a href="/admin">Return</a></div>')
        finally:connection.close()
        request.session.pop("uid",None);await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/sounds")
async def toggle_sounds(request:Request):
    if request.session.get("admin"):
        new="off" if sounds_enabled() else "on";connection=db();connection.execute("INSERT INTO settings(k,v) VALUES('sounds',?) ON CONFLICT(k) DO UPDATE SET v=excluded.v",(new,));connection.close();await hub.push()
    return RedirectResponse("/admin",303)


@app.post("/admin/backup")
def manual_backup(request:Request):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    path=backup_db("manual")
    return FileResponse(path,media_type="application/octet-stream",filename=path.name) if path else JSONResponse({"error":"No database found"},status_code=404)


def csv_response(rows,headers,name):
    output=io.StringIO();writer=csv.writer(output);writer.writerow(headers);writer.writerows(rows);return StreamingResponse(iter([output.getvalue()]),media_type="text/csv",headers={"Content-Disposition":f"attachment; filename={name}"})


@app.get("/admin/export/users")
def export_users(request:Request):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    connection=db();rows=connection.execute("SELECT name,email,balance FROM users WHERE mode=? ORDER BY name",(mode(),)).fetchall();connection.close();return csv_response(rows,["Name","Email","Points"],"participants.csv")


@app.get("/admin/export/results")
def export_results(request:Request):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    connection=db();rows=connection.execute("SELECT p.name,u.name,a.winning_bid,a.closed FROM auctions a JOIN prizes p ON p.id=a.prize_id LEFT JOIN users u ON u.id=a.winner WHERE a.mode=? AND a.status='CLOSED' ORDER BY a.id",(mode(),)).fetchall();connection.close();return csv_response(rows,["Prize","Winner","Winning Bid","Closed UTC"],"auction_results.csv")


@app.websocket("/ws")
async def websocket(websocket:WebSocket):
    await hub.add(websocket)
    try:
        while True:await websocket.receive_text()
    except WebSocketDisconnect:
        if websocket in hub.clients:hub.clients.remove(websocket)


@app.get("/api/display")
def display_state():
    active=current();winner=None
    if active and active["winner"]:
        connection=db()
        try:
            row=connection.execute("SELECT prize_name,winner_name,winning_bid FROM winners WHERE auction_id=?",(active["id"],)).fetchone()
            winner=dict(row) if row else None
        finally:connection.close()
    return JSONResponse({"auction":auction_payload(active),"server_now":time.time(),"winner":winner,"sounds":sounds_enabled()},headers={"Cache-Control":"no-store"})


def display_page():
    return page('''<div id="screen" class="display"><div id="trophy" class="trophy" hidden>🏆 WINNER 🏆</div><img id="image" class="prize" hidden><h2 id="status">Auction will begin soon</h2><div id="countdown" class="leader" role="timer"></div><h1 id="prize"></h1><p id="label"></p><div id="points" class="big"></div><div id="leader" class="leader"></div><small id="connection" role="status">Connecting...</small><button id="audio" class="secondary">Enable sounds</button></div><canvas id="confetti"></canvas>
    <script>
    const el=id=>document.getElementById(id);let previous=null,audio=null,soundOn=false,socket,retry,refreshing=false,pending=false,animation;
    el('audio').onclick=async()=>{try{audio=audio||new(window.AudioContext||window.webkitAudioContext)();await audio.resume();el('audio').hidden=true}catch(e){el('audio').textContent='Sounds unavailable'}};
    function tone(freq,duration){if(!soundOn||!audio||audio.state!=='running')return;const o=audio.createOscillator(),g=audio.createGain();o.frequency.value=freq;o.connect(g);g.connect(audio.destination);g.gain.setValueAtTime(.18,audio.currentTime);g.gain.exponentialRampToValueAtTime(.001,audio.currentTime+duration);o.start();o.stop(audio.currentTime+duration)}
    function confetti(){cancelAnimationFrame(animation);const c=el('confetti'),x=c.getContext('2d');c.width=innerWidth;c.height=innerHeight;let n=0,p=Array.from({length:180},()=>({x:Math.random()*c.width,y:-Math.random()*c.height,vx:(Math.random()-.5)*5,vy:2+Math.random()*5,h:Math.random()*360}));function frame(){x.clearRect(0,0,c.width,c.height);p.forEach(q=>{q.x+=q.vx;q.y+=q.vy;q.vy+=.03;x.fillStyle=`hsl(${q.h} 85% 55%)`;x.fillRect(q.x,q.y,6,6)});if(n++<240)animation=requestAnimationFrame(frame);else x.clearRect(0,0,c.width,c.height)}frame()}
    function render(s){const a=s.auction,w=s.winner,key=a?String(a.id):'none',won=!!(a&&a.status==='CLOSED'&&w);soundOn=s.sounds;updateAuctionClock(s);
      el('screen').classList.toggle('winner',won);el('trophy').hidden=!won;
      el('status').textContent=won?'Auction closed':auctionStatus(a);
      el('status').className=auctionIsOpen(a)?'open':'closed';el('prize').textContent=w?.prize_name||a?.prize||'';
      el('label').textContent=a?(won?'Winning bid':'Highest bid'):'';el('points').textContent=a?String(won?w.winning_bid:(a.high||0))+' points':'';
      el('leader').textContent=w?.winner_name||a?.leader||(a?'No bids yet':'');el('image').hidden=true;
      if(a?.image_url){const image=el('image');image.referrerPolicy='no-referrer';image.onerror=()=>{image.hidden=true};if(image.getAttribute('src')!==a.image_url)image.src=a.image_url;image.hidden=false}
      if(won&&(!previous||previous.key!==key||!previous.won)){confetti();tone(523,.25);setTimeout(()=>tone(659,.25),250);setTimeout(()=>tone(784,.45),500)}
      else if(a&&previous&&previous.key===key&&(a.high_bid_id||0)>previous.bid&&!won){el('points').classList.remove('bid-flash');void el('points').offsetWidth;el('points').classList.add('bid-flash');tone(880,.18)}
      if(previous&&previous.key!==key&&!won){cancelAnimationFrame(animation);el('confetti').getContext('2d').clearRect(0,0,innerWidth,innerHeight)}
      previous={key,won,bid:a?.high_bid_id||0};
    }
    async function refresh(){if(refreshing){pending=true;return}refreshing=true;try{const r=await fetch('/api/display',{cache:'no-store'});if(!r.ok)throw Error();const data=await r.json();updateAuctionClock(data);render(data)}catch(e){el('connection').textContent='Connection interrupted - retrying'}finally{refreshing=false;if(pending){pending=false;refresh()}}}
    function connect(){clearTimeout(retry);socket=new WebSocket((location.protocol==='https:'?'wss':'ws')+'://'+location.host+'/ws');socket.onopen=()=>{el('connection').textContent='Connected';refresh()};socket.onmessage=refresh;socket.onclose=()=>{el('connection').textContent='Reconnecting...';retry=setTimeout(connect,3000)};socket.onerror=()=>socket.close()}
    function resume(){refresh();if(!socket||socket.readyState===WebSocket.CLOSED)connect()}
    window.addEventListener('online',resume);document.addEventListener('visibilitychange',()=>{if(!document.hidden)resume()});refresh();connect();setInterval(()=>{if(!document.hidden)refresh()},1000);
    </script>''')


def participant_state(request):
    uid=request.session.get("uid");connection=db()
    try:
        connection.execute("BEGIN")
        m=connection.execute("SELECT v FROM settings WHERE k='mode'").fetchone()[0]
        user=session_user(request,connection,m)
        if not user:return None
        active=current(connection,m);a=auction_payload(active)
        personal=connection.execute("SELECT MAX(amount) FROM bids WHERE user_id=? AND auction_id=?",(uid,a["id"])).fetchone()[0] if a else None
        leading=connection.execute("SELECT user_id FROM bids WHERE id=?",(a["high_bid_id"],)).fetchone() if a and a["high_bid_id"] else None
        wins=[dict(x) for x in connection.execute("SELECT w.prize_name,w.winning_bid FROM winners w JOIN auctions a ON a.id=w.auction_id WHERE a.winner=? AND w.mode=? ORDER BY w.id DESC",(uid,m))]
        return {"server_now":time.time(),"name":user["name"],"balance":user["balance"],"auction":a,"last_bid":personal,"leading":bool(leading and leading[0]==uid),"wins":wins}
    finally:connection.close()


@app.get("/api/auction")
def auction_state(request:Request):
    state=participant_state(request)
    return JSONResponse(state if state else {"error":"Sign in again"},status_code=200 if state else 401,headers={"Cache-Control":"no-store"})


def participant_page(request):
    if not participant_state(request):return RedirectResponse("/",303)
    return page('''<div class="grid"><div class="card"><h2 id="rep-name"></h2><p>Remaining balance</p><div class="big" id="balance"></div><p id="connection" role="status">Connecting...</p><a href="/logout">Log out</a><h3>Your prizes</h3><p id="spent"></p><ul id="wins"></ul></div>
    <div class="card"><img id="prize-image" class="prize" hidden><h3 id="auction-status"></h3><div id="countdown" class="leader" role="timer"></div><h1 id="prize-name"></h1><p>Highest bid</p><div class="big" id="high"></div><div class="leader" id="leader"></div><p id="personal" role="status"></p>
    <form id="bid-form"><label for="amount">Your bid in points</label><input type="number" id="amount" min="1" required inputmode="numeric"><button id="submit">Submit Bid</button><button type="button" id="next">Bid 1 point more</button></form><p id="feedback" role="status" aria-live="polite"></p></div></div>
    <script>
    const el=id=>document.getElementById(id);let state,busy=false,socket,retry,refreshing=false,pending=false;
    function render(s){const a=s.auction,open=auctionIsOpen(a),changed=state&&((state.auction?.id||0)!==(a?.id||0));
      if(changed){el('amount').value='';el('feedback').textContent='The prize changed. Review it before bidding.'}
      state=s;el('rep-name').textContent=s.name;el('balance').textContent=s.balance;
      el('auction-status').textContent=auctionStatus(a);
      el('prize-name').textContent=a?.prize||'';el('high').textContent=a?.high||0;el('leader').textContent=a?.leader||'No bids yet';
      el('personal').textContent=s.last_bid?(open?(s.leading?'You are leading':'You have been outbid'):'Your last bid')+' - '+s.last_bid+' points':'You have not bid on this prize yet.';
      el('amount').max=s.balance;const minimum=(a?.high||0)+1;el('amount').min=minimum;
      el('submit').disabled=busy||!open;el('amount').disabled=busy||!open;el('next').disabled=busy||!open||minimum>s.balance;el('next').textContent='Bid '+minimum+' points';
      const image=el('prize-image');image.hidden=true;if(a?.image_url){image.referrerPolicy='no-referrer';image.onerror=()=>{image.hidden=true};if(image.getAttribute('src')!==a.image_url)image.src=a.image_url;image.hidden=false}
      el('wins').replaceChildren();for(const w of s.wins){const li=document.createElement('li');li.textContent=w.prize_name+' - '+w.winning_bid+' points';el('wins').append(li)}
      el('spent').textContent=s.wins.length?'Points spent: '+s.wins.reduce((n,w)=>n+w.winning_bid,0):'No prizes won yet.';
    }
    async function refresh(){if(refreshing){pending=true;return}refreshing=true;try{const r=await fetch('/api/auction',{cache:'no-store'});if(r.status===401){location.href='/';return}if(!r.ok)throw Error();const data=await r.json();updateAuctionClock(data);render(data)}catch(e){el('connection').textContent='Connection interrupted - retrying'}finally{refreshing=false;if(pending){pending=false;refresh()}}}
    function connect(){clearTimeout(retry);socket=new WebSocket((location.protocol==='https:'?'wss':'ws')+'://'+location.host+'/ws');socket.onopen=()=>{el('connection').textContent='Connected';refresh()};socket.onmessage=refresh;socket.onclose=()=>{el('connection').textContent='Reconnecting...';retry=setTimeout(connect,3000)};socket.onerror=()=>socket.close()}
    async function submit(amount){if(busy||!auctionIsOpen(state?.auction))return;const auction=state.auction.id;busy=true;render(state);el('feedback').textContent='Submitting...';try{const data=new URLSearchParams({amount,auction_id:auction});const r=await fetch('/bid',{method:'POST',headers:{Accept:'application/json'},body:data});const result=await r.json();el('feedback').textContent=result.message}catch(e){el('feedback').textContent='Could not confirm submission. Check your last bid before retrying.'}finally{busy=false;await refresh();if(state)render(state)}}
    el('bid-form').onsubmit=e=>{e.preventDefault();const amount=el('amount').value;if(confirm('Submit '+amount+' points for '+state.auction.prize+'?'))submit(amount)};
    el('next').onclick=()=>{const amount=(state.auction.high||0)+1;if(confirm('Submit '+amount+' points for '+state.auction.prize+'?'))submit(amount)};
    function heartbeat(){fetch('/heartbeat',{method:'POST'}).catch(()=>{})}
    window.addEventListener('online',()=>{refresh();if(!socket||socket.readyState===WebSocket.CLOSED)connect()});document.addEventListener('visibilitychange',()=>{if(!document.hidden){refresh();heartbeat();if(!socket||socket.readyState===WebSocket.CLOSED)connect()}});
    refresh();connect();heartbeat();setInterval(heartbeat,10000);setInterval(()=>{if(!document.hidden)refresh()},1000);
    </script>''')


@app.post("/admin/close/preview")
def preview_award(request:Request):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    a=current()
    if not a or a["status"]!="OPEN":return page('<h2>No open auction</h2><a href="/admin">Return</a>')
    if a['ends_at'] is not None and time.time() < a['ends_at']:
        return page('<h2>Bidding is still open</h2><p>Wait until the timer ends or use End Bidding Early to review the award.</p><a href="/admin">Return</a>')
    summary=f'Winner: {e(a["leader"])}<br>Winning bid: {a["high"]} points' if a["high_bid_id"] else 'No bids - close without awarding a prize.'
    return page(f'<div class="card"><h1>Confirm Award</h1><h2>{e(a["prize"])}</h2><p>{summary}</p><p>A new bid arriving before confirmation requires another review.</p><form method="post" action="/admin/close"><input type="hidden" name="auction_id" value="{a["id"]}"><input type="hidden" name="bid_id" value="{a["high_bid_id"] or 0}"><button class="danger">Confirm Close and Award</button></form><a href="/admin">Cancel</a></div>')


@app.post("/admin/import/preview")
async def preview_import(request:Request,workbook:UploadFile=File(...),action:str=Form("replace")):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    if action not in ("replace","append"):return JSONResponse({"error":"Invalid import action"},status_code=400)
    data=await workbook.read(10*1024*1024+1)
    if len(data)>10*1024*1024:return page('<h2>Workbook exceeds the 10 MB limit.</h2><a href="/admin">Return</a>')
    try:users,prizes,warnings=parse_xlsx(data)
    except Exception as error:return page(f'<h2>Preview failed</h2><p>{e(error)}</p><a href="/admin">Return</a>')
    if not users:return page('<h2>No valid participants found. Import cancelled.</h2><a href="/admin">Return</a>')
    folder=DATA_DIR/'import_previews';folder.mkdir(exist_ok=True)
    for old in folder.glob('*.xlsx'):
        if old.stat().st_mtime<time.time()-3600:old.unlink(missing_ok=True)
    token=secrets.token_hex(24);(folder/(token+'.xlsx')).write_bytes(data)
    request.session['import_preview']={"token":token,"mode":mode(),"action":action,"expires":time.time()+1800}
    warning_list=''.join(f'<li>{e(x)}</li>' for x in warnings)
    names=''.join(f'<li>{e(name)} - {e(email)} - {points} points</li>' for name,email,points in users[:10])
    prize_names=''.join(f'<li>{e(name)} - SKU {e(sku)} - {quantity} units</li>' for name,sku,image,quantity in prizes[:10])
    consequence='This deletes participants, prizes, bids, and winner history in the selected mode.' if action=='replace' else 'Existing balances and stock are preserved. Names and images update; new entries use workbook values.'
    return page(f'<div class="card"><h1>Import Preview - {mode().upper()}</h1><p>{len(users)} participants; {len(prizes)} unique prizes; {sum(p[3] for p in prizes)} prize units.</p><div class="warning">{e(consequence)}</div><h3>Warnings</h3><ul>{warning_list or "<li>None</li>"}</ul><h3>First 10 participants</h3><ul>{names}</ul><h3>First 10 prizes</h3><ul>{prize_names}</ul><form method="post" action="/admin/import/confirm"><input type="hidden" name="token" value="{token}"><button>Confirm Import</button></form><a href="/admin">Cancel</a></div>')


@app.post("/admin/import/confirm")
async def confirm_import(request:Request,token:str=Form(...)):
    if not request.session.get("admin"):return RedirectResponse("/admin",303)
    preview=request.session.get('import_preview')
    if not preview or not secrets.compare_digest(token,preview['token']) or preview['expires']<time.time() or preview['mode']!=mode():return page('<h2>Preview expired or mode changed. Preview again.</h2><a href="/admin">Return</a>')
    path=DATA_DIR/'import_previews'/(preview['token']+'.xlsx');claimed=path.with_suffix('.claimed')
    try:path.rename(claimed)
    except OSError:return page('<h2>This preview was already used. Preview again.</h2><a href="/admin">Return</a>')
    request.session.pop('import_preview',None)
    try:
        from starlette.datastructures import UploadFile as StoredUpload
        upload=StoredUpload(io.BytesIO(claimed.read_bytes()),filename='preview.xlsx')
        result=await import_workbook(request,upload,preview['action'],confirmed=True,expected_mode=preview['mode'])
        await hub.push()
        return result
    finally:claimed.unlink(missing_ok=True)


if __name__=="__main__":
    import uvicorn
    uvicorn.run("main:app",host="127.0.0.1",port=8000,proxy_headers=True,forwarded_allow_ips="127.0.0.1")
